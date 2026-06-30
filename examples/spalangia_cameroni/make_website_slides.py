#!/usr/bin/env python3
"""Spalangia cameroni - assembly + annotation summary (canonical v3).
Style: lowercase sentence text, neutral wording, tool/method shown per result.
Tool names and acronyms keep their canonical casing (they are identifiers)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY  = RGBColor(0x14, 0x2A, 0x4F)
BLUE  = RGBColor(0x1F, 0x6F, 0xB2)
TEAL  = RGBColor(0x12, 0x9A, 0x8F)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
AMBER = RGBColor(0xC8, 0x7B, 0x10)
RED   = RGBColor(0xC0, 0x40, 0x40)
GREY  = RGBColor(0x55, 0x5B, 0x66)
LGREY = RGBColor(0xEC, 0xEF, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK  = RGBColor(0x22, 0x27, 0x30)
SKY   = RGBColor(0xBF, 0xD8, 0xEE)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)

def box(s, l, t, w, h, fill=None, line=None, line_w=Pt(1)):
    sp = s.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = line_w
    return sp

def text(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=4):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sp_after); p.space_before = Pt(0)
        for (txt, size, color, bold) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = "Calibri"
    return tb

def header(s, title, kicker=None):
    box(s, 0, 0, 13.333, 1.15, fill=NAVY)
    box(s, 0, 1.15, 13.333, 0.06, fill=TEAL)
    text(s, 0.55, 0.16, 12.2, 0.95, [[(title, 29, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(s, 0.57, 0.68, 12.0, 0.4, [[(kicker, 12.5, SKY, False)]])

def footer(s, n):
    text(s, 0.55, 7.06, 9, 0.35, [[("Spalangia cameroni genome  ·  canonical v3  ·  2026-06-30", 9.5, GREY, False)]])
    text(s, 12.2, 7.06, 0.8, 0.35, [[(str(n), 9.5, GREY, False)]], align=PP_ALIGN.RIGHT)

def chip(s, l, t, w, h, big, small, src, accent):
    box(s, l, t, w, h, fill=WHITE, line=LGREY, line_w=Pt(1.25))
    box(s, l, t, 0.10, h, fill=accent)
    text(s, l+0.22, t+0.08, w-0.30, h-0.16,
         [[(big, 21, NAVY, True)], [(small, 10.5, GREY, False)], [(src, 9, accent, False)]],
         anchor=MSO_ANCHOR.MIDDLE, sp_after=1)

def bullets(s, l, t, w, h, items, color=DARK, size=13, gap=6):
    runs = []
    for it in items:
        if isinstance(it, tuple): txt, c, b = it
        else: txt, c, b = it, color, False
        runs.append([("–  ", size, TEAL, False), (txt, size, c, b)])
    text(s, l, t, w, h, runs, sp_after=gap)

def source(s, l, t, w, txt):
    text(s, l, t, w, 0.35, [[("source: ", 10, GREY, False), (txt, 10, GREY, False)]])

# =====================================================================
# 1 - TITLE
# =====================================================================
s = slide()
box(s, 0, 0, 13.333, 7.5, fill=NAVY)
box(s, 0, 4.55, 13.333, 0.07, fill=TEAL)
text(s, 0.9, 1.5, 11.5, 1.3, [[("Spalangia cameroni", 48, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.92, 2.85, 11.5, 0.7, [[("genome assembly & annotation — current status", 23, SKY, False)]])
text(s, 0.92, 3.55, 11.5, 0.6, [[("parasitoid wasp · PacBio HiFi + Illumina · RNA-Seq-guided annotation", 14, RGBColor(0x9A,0xB4,0xCF), False)]])
text(s, 0.92, 4.85, 11.6, 1.7,
     [[("canonical v3 gene set — assembly frozen, functional annotation complete", 14.5, WHITE, True)],
      [("assembly: 4,980 contigs · 651 Mb · busco 89.3%   (hifiasm + NextPolish; busco/miniprot)", 12.5, SKY, False)],
      [("annotation: 18,461 genes · 19,769 proteins · busco 83.0%   (BRAKER3/EVM/Tiberius/PASA)", 12.5, SKY, False)],
      [("function: PFAM + eggNOG + GO over the full proteome   (funannotate + emapper-2.1.13)", 12.5, SKY, False)]],
     sp_after=6)
text(s, 0.92, 6.75, 11.5, 0.4, [[("compiled 2026-06-30", 11.5, RGBColor(0x7E,0x9A,0xB8), False)]])

# =====================================================================
# 2 - INPUT DATA
# =====================================================================
s = slide(); header(s, "1 · input data & genome survey", "diploid female · PacBio Revio HiFi + Illumina NovaSeq X")
chip(s, 0.55, 1.5, 3.0, 1.35, "~12×", "HiFi coverage", "reads vs genome size", BLUE)
chip(s, 3.75, 1.5, 3.0, 1.35, "~80×", "Illumina coverage", "reads vs genome size", TEAL)
chip(s, 6.95, 1.5, 3.0, 1.35, "612 Mb", "haploid length", "GenomeScope 2.0, k=21", GREEN)
chip(s, 10.15, 1.5, 2.6, 1.35, "0.55%", "heterozygosity", "GenomeScope 2.0", AMBER)
text(s, 0.55, 3.15, 6.0, 0.4, [[("long reads — PacBio Revio HiFi", 14.5, NAVY, True)]])
bullets(s, 0.55, 3.6, 6.1, 1.7, [
    "866,363 reads · 7.88 Gb · read N50 9.6 kb · mean Q28",
    "sample GMCF_3514_04 (HiFi confirmed)",
])
source(s, 0.55, 4.75, 6.0, "NanoPlot / seqkit on raw HiFi")
text(s, 6.9, 3.15, 6.0, 0.4, [[("short reads — Illumina NovaSeq X", 14.5, NAVY, True)]])
bullets(s, 6.9, 3.6, 6.0, 1.7, [
    "363.6 M reads · 52.4 Gb · Q30 96.4% · GC 36.5%",
    "used for polishing, not contig construction",
])
source(s, 6.9, 4.75, 6.0, "fastp report (post-trim)")
box(s, 0.55, 5.45, 12.2, 1.15, fill=LGREY)
text(s, 0.8, 5.58, 11.8, 0.9,
     [[("repeat content estimated at ~30% of the genome (GenomeScope). coverage is 12× HiFi with no Hi-C, "
        "so the assembly is contig-level; this is stated as a limitation, not corrected for.", 12.5, DARK, False)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)

# =====================================================================
# 3 - ASSEMBLY WORKFLOW  (read-filter pre-step, 2 assemblers, Kaiju nr)
# =====================================================================
s = slide(); header(s, "2 · assembly workflow", "filter reads, then two assemblers compared, polish, decontaminate")
steps = [
    ("filter reads", "minimap2 map-hifi vs\ncontaminant panel", BLUE),
    ("assemble (2)", "hifiasm + Flye\ncompared", TEAL),
    ("select", "hifiasm primary\n(QUAST + busco)", GREEN),
    ("polish", "NextPolish\n(Illumina)", AMBER),
    ("decontaminate", "Kaiju vs NCBI nr\n+ insect rescue", RED),
    ("screen", "NCBI FCS-GX\n+ FCS-Adaptor", NAVY),
]
x = 0.55; w = 1.97
for i,(t1,t2,acc) in enumerate(steps):
    box(s, x, 1.65, w, 1.55, fill=WHITE, line=acc, line_w=Pt(2))
    box(s, x, 1.65, w, 0.42, fill=acc)
    text(s, x+0.06, 1.68, w-0.12, 0.4, [[(t1, 11.5, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x+0.08, 2.2, w-0.16, 0.95, [[(t2, 10, DARK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    if i < len(steps)-1:
        text(s, x+w-0.04, 1.95, 0.28, 0.5, [[("›", 20, GREY, True)]], anchor=MSO_ANCHOR.MIDDLE)
    x += w + 0.05
text(s, 0.55, 3.45, 12.0, 0.4, [[("two read-level / contig-level decontamination passes", 13.5, NAVY, True)]])
bullets(s, 0.55, 3.9, 12.2, 1.5, [
    "pre-assembly: HiFi reads screened with minimap2 (map-hifi) against a contaminant panel "
    "(human, PhiX, common bacteria including Wolbachia, fungi, vectors); reads mapping to it removed",
    "post-assembly: contigs classified with Kaiju against the NCBI nr protein database, "
    "with an insect-homology rescue step to avoid dropping real wasp contigs",
], gap=7)
chip(s, 0.55, 5.55, 3.0, 1.05, "4,980", "contigs", "QUAST / seqkit", BLUE)
chip(s, 3.75, 5.55, 3.0, 1.05, "651 Mb", "total length", "QUAST", TEAL)
chip(s, 6.95, 5.55, 3.0, 1.05, "273 kb", "contig N50", "QUAST", GREEN)
chip(s, 10.15, 5.55, 2.6, 1.05, "0 edits", "NCBI FCS screen", "FCS-GX + Adaptor", NAVY)
footer(s, 3)

# =====================================================================
# 4 - ASSEMBLY QC
# =====================================================================
s = slide(); header(s, "3 · assembly quality", "completeness and contiguity, with the method each came from")
text(s, 0.55, 1.5, 9, 0.4, [[("genome completeness", 14.5, NAVY, True)]])
bar_l, bar_t, bar_w, bar_h = 0.55, 1.98, 9.0, 0.55
box(s, bar_l, bar_t, bar_w, bar_h, fill=LGREY, line=GREY, line_w=Pt(0.75))
segs = [("S 77.4", 77.4, GREEN), ("D 11.9", 11.9, TEAL), ("F 1.9", 1.9, AMBER), ("M 8.8", 8.8, RED)]
xx = bar_l
for lab,pct,col in segs:
    sw = bar_w*pct/100.0
    box(s, xx, bar_t, sw, bar_h, fill=col)
    if pct > 4: text(s, xx, bar_t, sw, bar_h, [[(lab, 10.5, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    xx += sw
text(s, 0.55, 2.62, 9.0, 0.4, [[("complete 89.3%  (single 77.4 · duplicated 11.9) · fragmented 1.9 · missing 8.8", 12.5, DARK, False)]])
source(s, 0.55, 3.02, 9.0, "busco v6.0.0, hymenoptera_odb10 (n=5,991), euk_genome mode, miniprot predictor")
chip(s, 9.9, 1.5, 2.85, 0.92, "89.3%", "busco complete", "busco / miniprot", GREEN)
chip(s, 9.9, 2.52, 2.85, 0.92, "0 / 100kb", "ambiguous N's", "QUAST", BLUE)
text(s, 0.55, 3.55, 8, 0.4, [[("contiguity", 14.5, NAVY, True)]])
rows = [("metric","value"),("total length","650.9 Mb"),("# contigs","4,980"),
        ("largest contig","3.19 Mb"),("N50 / L50","273 kb / 631"),
        ("N90 / L90","50.4 kb / 2,687"),("GC content","37.06%")]
ty = 3.98; rh = 0.385
for i,(a,b) in enumerate(rows):
    fill = NAVY if i==0 else (WHITE if i%2 else LGREY)
    box(s, 0.55, ty, 6.2, rh, fill=fill)
    ca = WHITE if i==0 else DARK
    text(s, 0.7, ty, 3.8, rh, [[(a, 11.5, ca, i==0)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.4, ty, 2.2, rh, [[(b, 11.5, ca, i==0)]], anchor=MSO_ANCHOR.MIDDLE)
    ty += rh
source(s, 0.55, ty+0.02, 6.2, "QUAST on final_assembly.fa")
box(s, 7.1, 3.98, 5.65, 2.7, fill=LGREY)
text(s, 7.35, 4.15, 5.2, 2.4,
     [[("reading", 13.5, NAVY, True)],
      [("duplication 11.9% is consistent with residual haplotypic contigs (genome unphased, no Hi-C).", 11.5, DARK, False)],
      [("missing 8.8% is in line with the 12× HiFi coverage and ~30% repeat fraction.", 11.5, DARK, False)],
      [("figures are as reported by busco/QUAST; not adjusted.", 11.5, GREY, False)]],
     sp_after=8)
footer(s, 4)

# =====================================================================
# 5 - ANNOTATION WORKFLOW
# =====================================================================
s = slide(); header(s, "4 · annotation workflow", "RNA-Seq-guided, evidence-gated gene set")
text(s, 0.55, 1.42, 12, 0.4, [[("evidence inputs", 14.5, NAVY, True)]])
bullets(s, 0.55, 1.86, 6.2, 1.7, [
    "11 RNA-Seq libraries (public, venom, wholebody, NS5–8, S1–4)",
    "alignment with HISAT2/STAR → transcripts with StringTie",
    "genome repeat-masked with RepeatModeler/RepeatMasker",
    "protein homology from Nasonia via miniprot",
], gap=5)
text(s, 6.95, 1.42, 6, 0.4, [[("predictors combined", 14.5, NAVY, True)]])
bullets(s, 6.95, 1.86, 5.8, 1.7, [
    "BRAKER3 over all 11 BAMs (primary)",
    "EVidenceModeler consensus (core)",
    "Tiberius ab-initio (novel-locus rescue, GPU)",
    "miniprot busco-rescue at missing loci",
], gap=5)
chain = [
    ("BRAKER3 + EVM", BLUE),
    ("+ Tiberius rescue\n(573, StringTie-gated)", TEAL),
    ("+ PASA all-11\nUTRs / isoforms", GREEN),
    ("+ tRNAscan-SE (630)\n+ Rfam rRNA (8)", AMBER),
    ("canonical v3", NAVY),
]
text(s, 0.55, 3.95, 12, 0.4, [[("build chain (each step labelled with its tool)", 14.5, NAVY, True)]])
x=0.55; y=4.42
for i,(t1,acc) in enumerate(chain):
    w = 2.55 if i<4 else 1.9
    box(s, x, y, w, 1.0, fill=WHITE, line=acc, line_w=Pt(2))
    box(s, x, y, 0.09, 1.0, fill=acc)
    text(s, x+0.18, y+0.05, w-0.25, 0.9, [[(t1, 11, DARK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    x += w + 0.03
box(s, 0.55, 5.75, 12.2, 0.9, fill=LGREY)
text(s, 0.8, 5.86, 11.8, 0.7,
     [[("genes were added only at novel loci and only when supported by homology, RNA-Seq, or busco; "
        "existing models were not replaced blindly. the prior curated set was kept rather than rebuilt from scratch.", 12, DARK, False)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 5)

# =====================================================================
# 6 - ANNOTATION COMPLETENESS (busco per step vs genome + Nasonia)  [NEW]
# =====================================================================
s = slide(); header(s, "5 · annotation completeness across steps", "busco of each build step, against the genome and a reference proteome")
# embed the same figure used in the website bundle (02_busco_tracks.png) so the two match exactly.
# size by HEIGHT (the saved png's tight-bbox aspect ~1.689 differs from the raw figsize).
BUSCO_IMG = "/mnt/data/Projects/Elad_Chiel/wasp_genome_assembly/website_transfer/figures/02_busco_tracks.png"
pic_h = 4.85
pic_w = pic_h * 1.6894        # measured native aspect of the trimmed png
s.shapes.add_picture(BUSCO_IMG, Inches((13.333-pic_w)/2), Inches(1.42), height=Inches(pic_h))
text(s, 0.55, 1.42 + pic_h + 0.06, 12.2, 0.7,
     [[("completeness rises from 75% to 83% as evidence layers are added. the genome assembly itself sits at "
        "89.3% (a soft ceiling for annotation here), while a chromosome-scale relative, Nasonia vitripennis, "
        "reaches 96.3%. the remaining gap is attributed to coverage and masking, not predictor choice.", 11, DARK, False)]])
footer(s, 6)

# =====================================================================
# 7 - GENE SET (updated: pseudogenes + latest numbers)
# =====================================================================
s = slide(); header(s, "6 · canonical gene set (v3)", "structural annotation — final.gff3 / .gtf / .proteins.fa")
chip(s, 0.55, 1.5, 3.0, 1.3, "18,461", "genes", "final.gff3 counts", BLUE)
chip(s, 3.75, 1.5, 3.0, 1.3, "19,769", "mRNA / proteins", "final.gff3 counts", TEAL)
chip(s, 6.95, 1.5, 3.0, 1.3, "83.0%", "proteome busco complete", "busco, odb10, protein mode", GREEN)
chip(s, 10.15, 1.5, 2.6, 1.3, "300", "pseudogenes (flagged)", "internal-stop, 2026-06-30", AMBER)
text(s, 0.55, 3.1, 6, 0.4, [[("composition", 14.5, NAVY, True)]])
bullets(s, 0.55, 3.55, 6.4, 2.5, [
    "17,523 protein-coding genes (final.gff3)",
    "300 pseudogenes — sole-isoform internal-stop models, CDS removed",
    "630 tRNAs (tRNAscan-SE)",
    "8 rRNA loci (Infernal cmscan vs Rfam, StringTie-gated)",
    "573 Tiberius novel-locus rescues, flagged in col2",
], gap=5)
text(s, 7.1, 3.1, 5.6, 0.4, [[("UTR coverage", 14.5, NAVY, True)]])
bullets(s, 7.1, 3.55, 5.6, 2.5, [
    "12,937 five-prime UTRs",
    "10,653 three-prime UTRs",
    "~54% of mRNA carry a UTR",
    ("limited by short-read RNA only (no Iso-Seq)", GREY, False),
], gap=5)
source(s, 7.1, 5.35, 5.6, "PASA all-11 update on the gene set")
box(s, 0.55, 5.95, 12.2, 0.65, fill=LGREY)
text(s, 0.8, 6.0, 11.8, 0.55,
     [[("the 300 internal-stop models were flagged as pseudogenes (not deleted) and their CDS removed, so the "
        "proteome (19,769) has no internal stops. busco 83.0% was measured on the v3 proteome; flagging removed "
        "only broken models.", 11, DARK, False)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 7)

# =====================================================================
# 8 - FUNCTIONAL ANNOTATION
# =====================================================================
s = slide(); header(s, "7 · functional annotation", "run over the full v3 proteome, 2026-06-21")
chip(s, 0.55, 1.5, 2.95, 1.3, "11,825", "PFAM domains", "funannotate / hmmer3", BLUE)
chip(s, 3.70, 1.5, 2.95, 1.3, "14,178", "named products", "funannotate (UniProt+eggNOG)", TEAL)
chip(s, 6.85, 1.5, 2.95, 1.3, "9,486", "mRNA with GO terms", "eggNOG emapper-2.1.13", GREEN)
chip(s, 10.0, 1.5, 2.75, 1.3, "8,544", "eggNOG OGs", "emapper-2.1.13", AMBER)
text(s, 0.55, 3.1, 6, 0.4, [[("coverage", 14.5, NAVY, True)]])
bullets(s, 0.55, 3.55, 6.2, 2.3, [
    "COG 7,615 · KEGG_ko 8,625 (emapper-2.1.13)",
    "MEROPS proteases 642 · CAZymes 237 (funannotate)",
    "GO written as Ontology_term in final.gff3 (from eggNOG)",
    "covers BRK_* and tib_* genes, not only the evm core",
], gap=5)
text(s, 6.95, 3.1, 6, 0.4, [[("output files", 14.5, NAVY, True)]])
bullets(s, 6.95, 3.55, 5.8, 2.3, [
    "final.annotations.txt (eggNOG-enriched)",
    "final.eggnog.tsv · final.pfam.tsv",
    "final.gff3 (GO Ontology_term)",
    "proteins.fa (19,769; in-frame stops resolved)",
], gap=5)
box(s, 0.55, 5.95, 12.2, 0.65, fill=LGREY)
text(s, 0.8, 6.0, 11.8, 0.55,
     [[("InterProScan was not run; GO comes from eggNOG. the functional layer covers the whole current gene "
        "set, including the genes added by the BRAKER and Tiberius tracks.", 11.5, DARK, False)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 8)

# =====================================================================
# 9 - STATUS & NEXT (updated)
# =====================================================================
s = slide(); header(s, "8 · status & remaining work", "what is done and what is left, stated plainly")
text(s, 0.55, 1.45, 6, 0.4, [[("done", 15, GREEN, True)]])
bullets(s, 0.55, 1.95, 6.3, 3.4, [
    "assembly frozen; passes NCBI FCS with zero edits",
    "canonical v3 gene set staged (18,461 genes)",
    "300 internal-stop genes flagged as pseudogenes (CDS removed)",
    "proteome cleaned to 19,769 with no internal stops",
    "PASA all-11 UTRs staged (~54% of mRNA)",
    "functional layer complete (PFAM / eggNOG / GO)",
    "tRNA + rRNA grafted; website transfer bundle refreshed",
], gap=7)
text(s, 6.95, 1.45, 6, 0.4, [[("optional / deferred", 15, AMBER, True)]])
bullets(s, 6.95, 1.95, 5.8, 3.4, [
    "strip product names from the 300 pseudogenes (cosmetic)",
    "expression table (featureCounts) — PE/SE BAM mix to fix",
    "InterProScan domains (not needed for GO)",
    "Iso-Seq would raise UTR coverage above ~54%",
    ("Hi-C or deeper HiFi would address duplication / missing busco", GREY, False),
], gap=7)
box(s, 0.55, 5.6, 12.2, 1.0, fill=NAVY)
text(s, 0.8, 5.7, 11.8, 0.8,
     [[("the assembly is frozen and the annotation, including functional layers and pseudogene flagging, is "
        "complete. the items on the right are optional polish, not blockers.", 13.5, WHITE, False)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 9)

out = "/mnt/data/Projects/Elad_Chiel/wasp_genome_assembly/Spalangia_assembly_annotation_summary_20260630.pptx"
prs.save(out)
print("WROTE", out, "slides:", len(prs.slides._sldIdLst))
