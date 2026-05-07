#!/usr/bin/env python3
"""Add QC result slides to the S. cameroni genome assembly presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation("Spalangia_cameroni_assembly_plan.pptx")

# Color scheme (same as original)
DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MED_BLUE = RGBColor(0x2E, 0x6B, 0x9E)
LIGHT_BLUE = RGBColor(0x4A, 0x9E, 0xD9)
ACCENT_GREEN = RGBColor(0x27, 0xAE, 0x60)
ACCENT_ORANGE = RGBColor(0xE6, 0x7E, 0x22)
ACCENT_RED = RGBColor(0xC0, 0x39, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MED_GRAY = RGBColor(0x66, 0x66, 0x66)


def add_bg(slide, color=DARK_BLUE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=DARK_GRAY, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_slide(slide, left, top, width, height, items, font_size=16,
                     color=DARK_GRAY, spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = spacing
        p.level = 0
    return txBox


def add_rounded_rect(slide, left, top, width, height, fill_color, text="",
                     font_size=14, font_color=WHITE, bold=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(font_size)
        tf.paragraphs[0].font.color.rgb = font_color
        tf.paragraphs[0].font.bold = bold
        tf.paragraphs[0].font.name = "Calibri"
    return shape


def add_table(slide, left, top, width, height, rows, cols, data,
              header_color=MED_BLUE, header_font_color=WHITE):
    """Add a table with header row styled."""
    table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                                          Inches(width), Inches(height))
    table = table_shape.table

    for row_idx in range(rows):
        for col_idx in range(cols):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(data[row_idx][col_idx])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
                paragraph.font.name = "Calibri"
                if row_idx == 0:
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = header_font_color
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = header_color
                else:
                    paragraph.font.color.rgb = DARK_GRAY
                    if row_idx % 2 == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF5)
    return table_shape


# ============================================================
# SLIDE: Phase 1 QC Results — Section Divider
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BLUE)

add_text_box(slide, 1, 2.0, 11.3, 1.5,
             "Phase 1: QC Results",
             font_size=44, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 3.5, 11.3, 1.0,
             "PacBio HiFi & Illumina Short-Read Quality Assessment",
             font_size=24, color=LIGHT_BLUE, alignment=PP_ALIGN.CENTER)
add_text_box(slide, 1, 4.5, 11.3, 0.5,
             "Completed: April 15, 2026",
             font_size=16, color=RGBColor(0x88, 0xAA, 0xCC), alignment=PP_ALIGN.CENTER)


# ============================================================
# SLIDE: PacBio HiFi QC Results
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

add_text_box(slide, 0.5, 0.3, 12, 0.7,
             "PacBio HiFi Read QC Results (NanoPlot + seqkit)",
             font_size=30, color=DARK_BLUE, bold=True)

# Status badge
add_rounded_rect(slide, 10.5, 0.35, 2.3, 0.45, ACCENT_GREEN,
                 "CONFIRMED HiFi", 14, WHITE, True)

# Main stats table
pacbio_data = [
    ["Metric", "Value", "Assessment"],
    ["Read Type", "HiFi (CCS)", "Confirmed from /ccs headers"],
    ["Total Reads", "866,363", ""],
    ["Total Bases", "7.88 Gb", ""],
    ["Mean Read Length", "9,096 bp", "Typical HiFi range"],
    ["Median Read Length", "8,678 bp", ""],
    ["Read Length N50", "9,578 bp", "Good contiguity potential"],
    ["Longest Read", "35,538 bp", ""],
    ["Std Dev Length", "2,977 bp", "Tight distribution"],
    ["Mean Quality", "Q28.2", "Excellent HiFi quality"],
    ["Median Quality", "Q37.1", "Very high accuracy"],
    ["Q20 Reads", "96.1% (7,551 Mb)", ""],
    ["Q30 Reads", "76.1% (5,824 Mb)", "High-confidence bases"],
    ["GC Content", "36.73%", "Normal for Hymenoptera"],
]

add_table(slide, 0.5, 1.2, 8.0, 5.5, len(pacbio_data), 3, pacbio_data)

# Depth estimation box (right side)
add_rounded_rect(slide, 9.0, 1.2, 3.8, 0.45, DARK_BLUE,
                 "Coverage Depth Estimate", 14, WHITE, True)
add_bullet_slide(slide, 9.2, 1.8, 3.6, 3.5, [
    "Total HiFi bases: 7.88 Gb",
    "",
    "If genome ~200 Mb: ~39x",
    "If genome ~300 Mb: ~26x",
    "If genome ~400 Mb: ~20x",
    "",
    "Target: >=30x for hifiasm",
    "",
    "Depth will be refined after",
    "GenomeScope k-mer survey",
    "(Phase 1.3)",
], font_size=13, color=DARK_GRAY, spacing=Pt(2))

# Quality assessment box
add_rounded_rect(slide, 9.0, 5.0, 3.8, 0.45, ACCENT_GREEN,
                 "Quality: EXCELLENT", 14, WHITE, True)
add_bullet_slide(slide, 9.2, 5.5, 3.6, 1.5, [
    "Q30 > 76% of reads",
    "Tight length distribution",
    "No quality concerns",
    "Proceed with hifiasm",
], font_size=12, color=DARK_GRAY, spacing=Pt(2))


# ============================================================
# SLIDE: Illumina QC Results
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

add_text_box(slide, 0.5, 0.3, 12, 0.7,
             "Illumina NovaSeq X Read QC Results (fastp)",
             font_size=30, color=DARK_BLUE, bold=True)

add_rounded_rect(slide, 10.5, 0.35, 2.3, 0.45, ACCENT_GREEN,
                 "PASS", 14, WHITE, True)

# Before filtering
add_rounded_rect(slide, 0.5, 1.2, 6.0, 0.45, MED_BLUE,
                 "Before Filtering (Raw)", 14, WHITE, True)

before_data = [
    ["Metric", "Read 1", "Read 2"],
    ["Total Reads", "184,452,188", "184,452,188"],
    ["Total Bases", "27.67 Gb", "27.67 Gb"],
    ["Q20 (%)", "98.48%", "98.46%"],
    ["Q30 (%)", "94.70%", "94.97%"],
]
add_table(slide, 0.5, 1.8, 6.0, 2.2, len(before_data), 3, before_data)

# After filtering
add_rounded_rect(slide, 6.8, 1.2, 6.0, 0.45, ACCENT_GREEN,
                 "After Filtering (Trimmed)", 14, WHITE, True)

after_data = [
    ["Metric", "Read 1", "Read 2"],
    ["Total Reads", "182,165,699", "182,165,699"],
    ["Total Bases", "26.17 Gb", "26.17 Gb"],
    ["Q20 (%)", "99.19%", "99.29%"],
    ["Q30 (%)", "96.08%", "96.74%"],
]
add_table(slide, 6.8, 1.8, 6.0, 2.2, len(after_data), 3, after_data)

# Filtering summary
add_rounded_rect(slide, 0.5, 4.3, 12.3, 0.45, DARK_BLUE,
                 "Filtering Summary", 14, WHITE, True)

filter_data = [
    ["Metric", "Value", "Percentage"],
    ["Reads Passed Filter", "364,331,398", "98.76%"],
    ["Failed (Low Quality)", "3,874,494", "1.05%"],
    ["Failed (Too Many N)", "42,422", "0.01%"],
    ["Failed (Too Short)", "642,234", "0.17%"],
    ["Failed (Adapter Dimer)", "13,828", "<0.01%"],
    ["Reads with Adapter Trimmed", "73,732,466", "20.0%"],
    ["Bases Trimmed (Adapters)", "2.23 Gb", "4.0%"],
    ["Duplication Rate", "7.20%", "Low - good library"],
    ["Insert Size Peak", "181 bp", "Expected for short-insert lib"],
]
add_table(slide, 0.5, 4.9, 12.3, 3.8, len(filter_data), 3, filter_data)


# ============================================================
# SLIDE: QC Summary & Depth Analysis
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

add_text_box(slide, 0.5, 0.3, 12, 0.7,
             "Phase 1 QC Summary & Coverage Analysis",
             font_size=30, color=DARK_BLUE, bold=True)

# Combined summary
add_rounded_rect(slide, 0.5, 1.2, 6.0, 0.45, MED_BLUE,
                 "Data Summary", 14, WHITE, True)

summary_data = [
    ["Data Type", "Total Reads", "Total Bases", "Quality"],
    ["PacBio HiFi", "866,363", "7.88 Gb", "Q28.2 mean"],
    ["Illumina (trimmed)", "364,331,398", "52.33 Gb", "Q30 >96%"],
    ["Combined", "365,197,761", "60.21 Gb", ""],
]
add_table(slide, 0.5, 1.8, 6.0, 1.8, len(summary_data), 4, summary_data)

# Coverage estimates
add_rounded_rect(slide, 6.8, 1.2, 6.0, 0.45, ACCENT_ORANGE,
                 "Estimated Coverage (pending GenomeScope)", 14, WHITE, True)

cov_data = [
    ["Est. Genome Size", "PacBio HiFi", "Illumina", "Total"],
    ["200 Mb", "~39x", "~262x", "~301x"],
    ["300 Mb", "~26x", "~174x", "~201x"],
    ["400 Mb", "~20x", "~131x", "~151x"],
]
add_table(slide, 6.8, 1.8, 6.0, 1.8, len(cov_data), 4, cov_data)

# Key findings
add_rounded_rect(slide, 0.5, 4.0, 6.0, 0.45, ACCENT_GREEN,
                 "Key Findings", 14, WHITE, True)
add_bullet_slide(slide, 0.7, 4.55, 5.8, 2.5, [
    "PacBio reads confirmed as HiFi (CCS) from Revio platform",
    "HiFi quality excellent: median Q37.1, 76% reads >Q30",
    "Read N50 = 9,578 bp — good for resolving repeats",
    "Illumina data high quality: 98.76% reads pass filtering",
    "Low duplication rate (7.2%) — good library complexity",
    "Insert size 181 bp — consistent with short-insert library",
    "GC content 36.73% — normal for Hymenoptera",
], font_size=14, color=DARK_GRAY, spacing=Pt(3))

# Action items
add_rounded_rect(slide, 6.8, 4.0, 6.0, 0.45, DARK_BLUE,
                 "Next Steps", 14, WHITE, True)
add_bullet_slide(slide, 7.0, 4.55, 5.8, 2.5, [
    "Phase 1.3: K-mer survey (Jellyfish + GenomeScope 2.0)",
    "  -> Determine exact genome size & heterozygosity",
    "  -> Refine coverage depth estimates",
    "  -> Assess if sample is haploid male or diploid female",
    "",
    "Phase 2.1: Contig assembly (hifiasm)",
    "  -> Use confirmed HiFi reads",
    "  -> Expect contig N50 in Mb range",
], font_size=14, color=DARK_GRAY, spacing=Pt(3))

# Status bar
add_rounded_rect(slide, 0.5, 6.8, 12.3, 0.5, ACCENT_GREEN,
                 "STATUS: Phase 1.1 & 1.2 COMPLETE  |  Phase 1.3 (k-mer survey) NEXT",
                 15, WHITE, True)


# Save
output_path = "Spalangia_cameroni_assembly_plan.pptx"
prs.save(output_path)
print(f"Presentation updated with QC slides: {output_path}")
print(f"Total slides: {len(prs.slides)}")
