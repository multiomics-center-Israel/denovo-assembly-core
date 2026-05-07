#!/usr/bin/env python3
"""
Add coverage analysis and dual-mode assembly strategy slides
to the Spalangia cameroni assembly plan presentation.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Colors ──────────────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x1B, 0x3A, 0x5C)
MED_BLUE    = RGBColor(0x2E, 0x6B, 0x9E)
LIGHT_BLUE  = RGBColor(0x4A, 0x9E, 0xD9)
ACCENT_GREEN  = RGBColor(0x27, 0xAE, 0x60)
ACCENT_ORANGE = RGBColor(0xE6, 0x7E, 0x22)
ACCENT_RED    = RGBColor(0xC0, 0x39, 0x2B)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY   = RGBColor(0x33, 0x33, 0x33)
MED_GRAY    = RGBColor(0x66, 0x66, 0x66)

FONT = "Calibri"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# ── Helpers ─────────────────────────────────────────────────────────

def _set_font(run, size=12, bold=False, color=DARK_GRAY, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic


def _add_para(tf, text, size=12, bold=False, color=DARK_GRAY,
              alignment=PP_ALIGN.LEFT, space_before=Pt(2), space_after=Pt(2),
              bullet=False, italic=False, level=0):
    p = tf.add_paragraph()
    p.alignment = alignment
    p.space_before = space_before
    p.space_after = space_after
    p.level = level
    if bullet:
        p.level = level if level else 0
    run = p.add_run()
    run.text = text
    _set_font(run, size, bold, color, italic)
    return p


def _rounded_rect(slide, left, top, width, height, fill_color,
                  border_color=None, border_width=Pt(1)):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    return shape


def _header_bar(slide, left, top, width, height, text, fill_color,
                font_size=14, font_color=WHITE):
    shape = _rounded_rect(slide, left, top, width, height, fill_color)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = text
    _set_font(run, font_size, bold=True, color=font_color)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    return shape


def _badge(slide, left, top, width, height, text, fill_color,
           font_size=10, font_color=WHITE):
    return _header_bar(slide, left, top, width, height, text,
                       fill_color, font_size, font_color)


def _content_box(slide, left, top, width, height, border_color=MED_BLUE):
    shape = _rounded_rect(slide, left, top, width, height,
                          WHITE, border_color, Pt(1.5))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(6)
    tf.margin_bottom = Pt(6)
    return shape, tf


def _slide_title(slide, text):
    """Dark blue title bar across top."""
    bar = _rounded_rect(slide, Inches(0.3), Inches(0.25),
                        Inches(12.7), Inches(0.7), DARK_BLUE)
    tf = bar.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    tf.margin_left = Pt(14)
    run = tf.paragraphs[0].add_run()
    run.text = text
    _set_font(run, 26, bold=True, color=WHITE)
    return bar


def _add_table(slide, rows, cols, left, top, width, height):
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    return table_shape.table


def _style_table_cell(cell, text, size=10, bold=False, color=DARK_GRAY,
                      fill=None, alignment=PP_ALIGN.LEFT):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    _set_font(run, size, bold, color)
    cell.margin_left = Pt(5)
    cell.margin_right = Pt(5)
    cell.margin_top = Pt(3)
    cell.margin_bottom = Pt(3)
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


# ════════════════════════════════════════════════════════════════════
#  SLIDE A – Coverage Analysis & Critical Decision Point
# ════════════════════════════════════════════════════════════════════

def add_slide_a(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank

    _slide_title(slide, "Coverage Analysis & Assembly Strategy Decision")

    # ── Left box: Genome Size Discovery ─────────────────────────
    left_x = Inches(0.3)
    box_top = Inches(1.15)
    box_w = Inches(6.1)

    _header_bar(slide, left_x, box_top, box_w, Inches(0.4),
                "Genome Size Discovery", MED_BLUE, 13)

    shape_l, tf_l = _content_box(slide, left_x, Inches(1.6), box_w, Inches(3.0))
    # Clear default paragraph
    tf_l.paragraphs[0].clear()

    bullets = [
        ("Expected genome size: ", "200-400 Mb", " (based on typical Chalcidoidea)"),
        ("Actual genome size: ", "~655 Mb", " (GenomeScope 2.0)"),
        ("This is ", "~2x larger", " than expected"),
        ("Comparable to ", "Eretmocerus hayati (692 Mb)", " — likely TE expansion"),
        ("35% repeat content ", "(~229 Mb)", ""),
        ("Heterozygosity: ", "0.55%", " — confirms diploid female"),
    ]
    for i, (pre, highlight, post) in enumerate(bullets):
        if i == 0:
            p = tf_l.paragraphs[0]
        else:
            p = tf_l.add_paragraph()
        p.space_before = Pt(3)
        p.space_after = Pt(3)
        p.level = 0

        r1 = p.add_run()
        r1.text = "  \u2022  " + pre
        _set_font(r1, 11, color=DARK_GRAY)

        r2 = p.add_run()
        r2.text = highlight
        _set_font(r2, 11, bold=True, color=DARK_BLUE)

        if post:
            r3 = p.add_run()
            r3.text = post
            _set_font(r3, 11, color=DARK_GRAY)

    # ── Right box: Coverage Impact ──────────────────────────────
    right_x = Inches(6.7)
    right_w = Inches(6.3)

    _header_bar(slide, right_x, box_top, Inches(4.5), Inches(0.4),
                "Coverage Impact", MED_BLUE, 13)
    # RED WARNING badge
    _badge(slide, Inches(11.4), box_top, Inches(1.6), Inches(0.4),
           "\u26A0  WARNING", ACCENT_RED, 11)

    # Coverage table
    tbl = _add_table(slide, 6, 3, right_x, Inches(1.65), right_w, Inches(2.6))

    headers = ["Metric", "Value", "Assessment"]
    data = [
        ("Genome size", "~655 Mb", "2x larger than expected"),
        ("PacBio HiFi total", "7.88 Gb", ""),
        ("PacBio coverage", "~12x", "CRITICAL: Below 30x minimum"),
        ("Illumina total", "52.33 Gb", ""),
        ("Illumina coverage", "~80x", "Sufficient for polishing"),
    ]

    col_widths = [Inches(1.8), Inches(1.5), Inches(3.0)]
    for ci, w in enumerate(col_widths):
        tbl.columns[ci].width = w

    for ci, h in enumerate(headers):
        _style_table_cell(tbl.cell(0, ci), h, 10, True, WHITE,
                          fill=DARK_BLUE, alignment=PP_ALIGN.CENTER)

    for ri, (m, v, a) in enumerate(data, start=1):
        _style_table_cell(tbl.cell(ri, 0), m, 10, False, DARK_GRAY)
        _style_table_cell(tbl.cell(ri, 1), v, 10, True, DARK_BLUE,
                          alignment=PP_ALIGN.CENTER)
        color = ACCENT_RED if "CRITICAL" in a else DARK_GRAY
        _style_table_cell(tbl.cell(ri, 2), a, 10, "CRITICAL" in a, color)
        # alternate row shading
        if ri % 2 == 0:
            for ci in range(3):
                cell = tbl.cell(ri, ci)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xF8)

    # ── Bottom: Implications ────────────────────────────────────
    impl_top = Inches(4.85)
    _header_bar(slide, Inches(0.3), impl_top, Inches(12.7), Inches(0.4),
                "Implications for Assembly", ACCENT_ORANGE, 13)

    shape_b, tf_b = _content_box(slide, Inches(0.3), Inches(5.3),
                                 Inches(12.7), Inches(1.9))
    tf_b.paragraphs[0].clear()

    implications = [
        "Assembly will be more fragmented (lower N50)",
        "Repetitive regions (35%) harder to resolve at 12x",
        "Some genes may be incomplete or split across contigs",
        "Haplotig purging essential (diploid female)",
    ]
    for i, txt in enumerate(implications):
        if i == 0:
            p = tf_b.paragraphs[0]
        else:
            p = tf_b.add_paragraph()
        p.space_before = Pt(4)
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = "  \u2022  " + txt
        _set_font(r, 12, color=DARK_GRAY)


# ════════════════════════════════════════════════════════════════════
#  SLIDE B – Dual-Mode Assembly Strategy
# ════════════════════════════════════════════════════════════════════

def add_slide_b(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_title(slide, "Dual-Mode Assembly Strategy")

    col_w = Inches(6.1)

    # ── LEFT column: Mode 1 ─────────────────────────────────────
    lx = Inches(0.3)
    _header_bar(slide, lx, Inches(1.15), col_w, Inches(0.45),
                "Mode 1: HiFi-Only (hifiasm)", MED_BLUE, 14)

    shape_l, tf_l = _content_box(slide, lx, Inches(1.65), col_w, Inches(4.3))
    tf_l.paragraphs[0].clear()

    mode1 = [
        ("Assembler: ", "hifiasm"),
        ("Input: ", "PacBio HiFi reads only (12x)"),
        ("Pros: ", "Gold-standard for HiFi, handles heterozygosity natively, proven pipeline"),
        ("Cons: ", "12x is below recommended 30x, may produce fragmented assembly"),
        ("Expected N50: ", "Lower than typical (possibly 100 kb - 1 Mb range)"),
        ("Polish: ", "NextPolish with Illumina (2 rounds)"),
        ("Best case: ", "Usable draft assembly with BUSCO >90%"),
        ("When to use: ", "Standard approach, compare results with Mode 2"),
    ]
    for i, (label, value) in enumerate(mode1):
        p = tf_l.paragraphs[0] if i == 0 else tf_l.add_paragraph()
        p.space_before = Pt(4)
        p.space_after = Pt(4)
        rl = p.add_run()
        rl.text = "  \u2022  " + label
        _set_font(rl, 11, bold=True, color=DARK_BLUE)
        rv = p.add_run()
        rv.text = value
        _set_font(rv, 11, color=DARK_GRAY)

    # ── RIGHT column: Mode 2 ────────────────────────────────────
    rx = Inches(6.7)
    _header_bar(slide, rx, Inches(1.15), col_w, Inches(0.45),
                "Mode 2: Hybrid HiFi+Illumina (MaSuRCA)", ACCENT_GREEN, 14)

    shape_r, tf_r = _content_box(slide, rx, Inches(1.65), col_w, Inches(4.3))
    tf_r.paragraphs[0].clear()

    mode2 = [
        ("Assembler: ", "MaSuRCA (Maryland Super-Read Celera Assembler)"),
        ("Input: ", "PacBio HiFi + Illumina reads combined"),
        ("Pros: ", "Compensates for low long-read depth using Illumina, designed for hybrid assembly, may produce higher contiguity"),
        ("Cons: ", "More complex, longer runtime, less established for HiFi"),
        ("Expected N50: ", "Potentially higher than Mode 1"),
        ("Polish: ", "Built-in polishing"),
        ("Best case: ", "Better contiguity by leveraging 80x Illumina depth"),
        ("When to use: ", "When long-read coverage is insufficient"),
    ]
    for i, (label, value) in enumerate(mode2):
        p = tf_r.paragraphs[0] if i == 0 else tf_r.add_paragraph()
        p.space_before = Pt(4)
        p.space_after = Pt(4)
        rl = p.add_run()
        rl.text = "  \u2022  " + label
        _set_font(rl, 11, bold=True, color=ACCENT_GREEN if True else DARK_BLUE)
        rv = p.add_run()
        rv.text = value
        _set_font(rv, 11, color=DARK_GRAY)

    # ── Bottom: Recommended Approach ────────────────────────────
    rec_shape = _rounded_rect(slide, Inches(0.3), Inches(6.15),
                              Inches(12.7), Inches(1.05), DARK_BLUE)
    tf_rec = rec_shape.text_frame
    tf_rec.word_wrap = True
    tf_rec.margin_left = Pt(14)
    tf_rec.margin_right = Pt(14)
    tf_rec.margin_top = Pt(8)

    p0 = tf_rec.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    r0 = p0.add_run()
    r0.text = "Recommended Approach"
    _set_font(r0, 13, bold=True, color=WHITE)

    p1 = tf_rec.add_paragraph()
    p1.alignment = PP_ALIGN.LEFT
    p1.space_before = Pt(4)
    r1 = p1.add_run()
    r1.text = ("Run both modes in parallel, compare BUSCO, N50, and total assembly size. "
               "Select the better assembly for downstream annotation. If neither meets "
               "quality threshold, request additional PacBio sequencing.")
    _set_font(r1, 11, color=WHITE)


# ════════════════════════════════════════════════════════════════════
#  SLIDE C – Options for Improving Assembly Quality
# ════════════════════════════════════════════════════════════════════

def add_slide_c(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_title(slide, "Options for Improving Assembly Quality")

    col_w = Inches(4.0)
    col_gap = Inches(0.15)
    start_x = Inches(0.3)
    badge_h = Inches(0.4)
    box_top = Inches(1.6)
    box_h = Inches(4.6)

    # ── Column 1: Option A ──────────────────────────────────────
    x1 = start_x
    _badge(slide, x1, Inches(1.15), col_w, badge_h,
           "Option A: Proceed As-Is", ACCENT_ORANGE, 13)

    sh1, tf1 = _content_box(slide, x1, box_top, col_w, box_h)
    tf1.paragraphs[0].clear()
    optA = [
        "Run both assembly modes",
        "Compare results",
        "If BUSCO >90%, proceed to annotation",
        "Fastest path to results",
        "Risk: fragmented assembly",
    ]
    for i, txt in enumerate(optA):
        p = tf1.paragraphs[0] if i == 0 else tf1.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = "  \u2022  " + txt
        bold = "Risk" in txt
        color = ACCENT_RED if "Risk" in txt else DARK_GRAY
        _set_font(r, 12, bold=bold, color=color)

    # ── Column 2: Option B ──────────────────────────────────────
    x2 = start_x + col_w + col_gap
    _badge(slide, x2, Inches(1.15), col_w, badge_h,
           "Option B: Additional PacBio Sequencing", ACCENT_GREEN, 12)
    # "RECOMMENDED" sub-badge
    _badge(slide, x2 + col_w - Inches(1.6), Inches(1.15), Inches(1.6), badge_h,
           "RECOMMENDED", ACCENT_GREEN, 10)

    sh2, tf2 = _content_box(slide, x2, box_top, col_w, box_h,
                            border_color=ACCENT_GREEN)
    tf2.paragraphs[0].clear()
    optB = [
        "Request 1 additional Revio SMRT cell",
        "Would add ~8 Gb \u2192 total ~16 Gb \u2192 ~24x coverage",
        "Much closer to 30x target",
        "Significant improvement in contiguity expected",
        "Timeline: depends on facility availability",
    ]
    for i, txt in enumerate(optB):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = "  \u2022  " + txt
        _set_font(r, 12, color=DARK_GRAY)

    # ── Column 3: Option C ──────────────────────────────────────
    x3 = start_x + 2 * (col_w + col_gap)
    _badge(slide, x3, Inches(1.15), col_w, badge_h,
           "Option C: Alternative Long-Read Data", LIGHT_BLUE, 12)

    sh3, tf3 = _content_box(slide, x3, box_top, col_w, box_h,
                            border_color=LIGHT_BLUE)
    tf3.paragraphs[0].clear()
    optC = [
        "Consider Oxford Nanopore (ONT) ultra-long reads",
        "Complement HiFi with ONT for scaffolding",
        "Tools: hifiasm can use both HiFi + ONT",
        "Would help span repetitive regions",
    ]
    for i, txt in enumerate(optC):
        p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = "  \u2022  " + txt
        _set_font(r, 12, color=DARK_GRAY)

    # ── Bottom status bar ───────────────────────────────────────
    status = _rounded_rect(slide, Inches(0.3), Inches(6.45),
                           Inches(12.7), Inches(0.75), DARK_BLUE)
    tf_s = status.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = Pt(14)
    tf_s.margin_top = Pt(8)
    p = tf_s.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER

    r1 = p.add_run()
    r1.text = "DECISION: "
    _set_font(r1, 14, bold=True, color=WHITE)

    r2 = p.add_run()
    r2.text = ("Proceeding with dual-mode assembly (Option A). "
               "Additional sequencing (Option B) under investigation.")
    _set_font(r2, 13, color=WHITE)


# ════════════════════════════════════════════════════════════════════
#  Main
# ════════════════════════════════════════════════════════════════════

def main():
    pptx_path = "Spalangia_cameroni_assembly_plan.pptx"
    prs = Presentation(pptx_path)

    add_slide_a(prs)
    add_slide_b(prs)
    add_slide_c(prs)

    prs.save(pptx_path)
    print(f"Done — saved {pptx_path} with {len(prs.slides)} slides "
          f"(added 3 new slides: slides 18-20)")


if __name__ == "__main__":
    main()
