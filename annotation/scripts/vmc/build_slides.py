#!/usr/bin/env python3
"""Build an updated results deck from the current numbers + figures.
Deterministic; reads the TSVs the driver just wrote. No network.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

PROJ = "/mnt/data/Projects/Elad_Chiel/wasp_genome_assembly"
OUT = os.path.join(PROJ, "report/Spalangia_results_update.pptx")
FIG = os.path.join(PROJ, "annotation/figures")

def read_tsv(path):
    if not os.path.exists(path):
        return []
    with open(path) as fh:
        return [l.rstrip("\n").split("\t") for l in fh if l.strip()]

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]; TITLE = prs.slide_layouts[0]

def add_title(slide, text, size=30):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.9))
    p = tb.text_frame.paragraphs[0]; p.text = text
    p.font.size = Pt(size); p.font.bold = True; p.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)

def add_bullets(slide, lines, top=1.3, size=16, left=0.6, width=12.0):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.8))
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        lvl = 0
        while ln.startswith("  "):
            ln = ln[2:]; lvl += 1
        p.text = ln; p.level = lvl
        p.font.size = Pt(size - 2*lvl)

def add_image(slide, path, left, top, width=None, height=None):
    if os.path.exists(path):
        kw = {}
        if width: kw["width"] = Inches(width)
        if height: kw["height"] = Inches(height)
        slide.shapes.add_picture(path, Inches(left), Inches(top), **kw)

def add_table(slide, rows, left=0.5, top=1.4, width=12.3, height=5.5, fs=11):
    if not rows:
        return
    nr, nc = len(rows), max(len(r) for r in rows)
    t = slide.shapes.add_table(nr, nc, Inches(left), Inches(top),
                               Inches(width), Inches(height)).table
    for ci in range(nc):
        for ri, row in enumerate(rows):
            cell = t.cell(ri, ci)
            cell.text = row[ci] if ci < len(row) else ""
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(fs)
                if ri == 0: p.font.bold = True

# --- Title slide ---
s = prs.slides.add_slide(TITLE)
s.shapes.title.text = "Spalangia cameroni — Genome & Annotation Update"
s.placeholders[1].text = ("Quad-mode assembly · EVM+PASA canonical annotation · "
                          "venom & mitochondrial survey · comparison to Nasonia vitripennis")

# --- Pipeline diagram ---
s = prs.slides.add_slide(BLANK)
add_title(s, "Full pipeline — methods overview")
add_image(s, os.path.join(FIG, "pipeline_diagram.png"), 0.4, 1.2, width=12.5)

# --- Cross-stage comparison table ---
s = prs.slides.add_slide(BLANK)
add_title(s, "Assembly & annotation quality across stages (vs Nasonia)")
rows = read_tsv(os.path.join(PROJ, "analysis/comparison/annotation_comparison.tsv"))
add_table(s, rows, fs=10)

# --- BUSCO figure ---
s = prs.slides.add_slide(BLANK)
add_title(s, "BUSCO completeness across annotation stages")
add_image(s, os.path.join(FIG, "stage_comparison.png"), 1.2, 1.2, width=10.8)

# --- Canonical set summary ---
s = prs.slides.add_slide(BLANK)
add_title(s, "Canonical annotation: EVM+PASA + funannotate")
add_bullets(s, [
    "Canonical gene set = EVM consensus (BRAKER + rescued + PASA + Nasonia homology), PASA-updated, funannotate-annotated",
    "12,580 gene models; BUSCO(protein) C:75.1% [S:65.9%, D:9.1%], F:3.8%, M:21.1%",
    "Duplication dropped to 9.1% (vs 21.6% in the RNA-seq 'rescued' set) — cleaner, NCBI-ready structures",
    "Functional layers: PFAM (16,312), UniProt names (1,087), MEROPS proteases (484), dbCAN CAZymes (214), BUSCO models (3,860)",
    "Gaps to flag: eggNOG + InterProScan not run; secretome (SignalP) pending → functional naming relies on Nasonia transfer + PFAM",
])

# --- Venom ---
s = prs.slides.add_slide(BLANK)
add_title(s, "Venom-gland gene recovery")
vrows = read_tsv(os.path.join(PROJ, "analysis/venom/venom_recovery_summary.tsv"))
add_table(s, vrows, top=1.4, height=2.5, fs=12)
add_bullets(s, [
    "Reference = Nasonia/parasitoid venom proteins (UniProt).",
    "Caveat: only one public RNA-seq library and it is whole-body (not venom gland) — recovery rests on protein homology + ab initio, not expression.",
], top=4.2, size=14)

# --- Mito ---
s = prs.slides.add_slide(BLANK)
add_title(s, "Mitochondrial contig")
mrows = read_tsv(os.path.join(PROJ, "analysis/mito/mito_summary.tsv"))
add_table(s, [r[:4] for r in mrows], top=1.4, height=3.5, fs=11)

prs.save(OUT)
print("wrote", OUT)
